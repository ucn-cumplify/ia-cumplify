"""Reads the text out of an evidence file.

The file comes back as locatable fragments: one per PDF page, per Excel sheet and per
PowerPoint slide, and a single body for Word (paragraphs and tables in document order)
and for CSV. Each file also comes back with a status: read, unreadable (a scanned PDF
or an image), unsupported or error. Size limits and truncation are not applied here.
"""

from __future__ import annotations

import csv
from collections.abc import Callable
from datetime import date, datetime
from io import BytesIO, StringIO

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from ia_cumplify.domain.evidence_text import (
    BODY,
    ERROR,
    PAGE,
    READ,
    SHEET,
    SLIDE,
    UNREADABLE,
    UNSUPPORTED,
    ExtractedFile,
    ExtractedFragment,
)

# Raster images have no text layer. They are not parsed.
_IMAGE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"})

_READERS: dict[str, Callable[[bytes], tuple[ExtractedFragment, ...]]] = {}


def extract_text(extension: str, content: bytes) -> ExtractedFile:
    """Read ``content``. ``extension`` includes the dot, as in ``.pdf``.

    A reader that raises becomes ``error`` and does not propagate. An extension
    with no reader is ``unsupported``. An image, or a file whose fragments are
    all blank, is ``unreadable``.
    """
    normalized = extension.strip().lower()
    if normalized and not normalized.startswith("."):
        normalized = f".{normalized}"

    if normalized in _IMAGE_EXTENSIONS:
        return ExtractedFile(UNREADABLE, ())

    reader = _READERS.get(normalized)
    if reader is None:
        return ExtractedFile(UNSUPPORTED, ())

    try:
        fragments = reader(content)
    except Exception as exc:
        return ExtractedFile(ERROR, (), _error_detail(exc))

    if not any(fragment.text.strip() for fragment in fragments):
        return ExtractedFile(UNREADABLE, fragments)
    return ExtractedFile(READ, fragments)


def _error_detail(exc: Exception) -> str:
    message = " ".join(str(exc).split())
    if not message:
        return type(exc).__name__
    return message[:500]


def extract_pdf(content: bytes) -> tuple[ExtractedFragment, ...]:
    reader = PdfReader(BytesIO(content))
    return tuple(
        ExtractedFragment(PAGE, str(number), (page.extract_text() or "").strip())
        for number, page in enumerate(reader.pages, start=1)
    )


def extract_docx(content: bytes) -> tuple[ExtractedFragment, ...]:
    document = Document(BytesIO(content))
    parts: list[str] = []
    for block in _iter_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                parts.append(text)
        else:
            rendered = _render_table(block)
            if rendered:
                parts.append(rendered)
    return (ExtractedFragment(BODY, "", "\n".join(parts)),)


def extract_xlsx(content: bytes) -> tuple[ExtractedFragment, ...]:
    # read_only walks rows without building the full cell graph. data_only keeps the
    # stored value; a formula without a cached result comes back empty.
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        return tuple(
            ExtractedFragment(SHEET, sheet.title, _render_rows(sheet.iter_rows(values_only=True)))
            for sheet in workbook.worksheets
        )
    finally:
        workbook.close()


def extract_csv(content: bytes) -> tuple[ExtractedFragment, ...]:
    text = _decode_csv(content)
    delimiter = _detect_delimiter(text)
    rows = csv.reader(StringIO(text), delimiter=delimiter)
    rendered = "\n".join(
        delimiter.join(cell.strip() for cell in row)
        for row in rows
        if any(cell.strip() for cell in row)
    )
    return (ExtractedFragment(BODY, "", rendered),)


def extract_pptx(content: bytes) -> tuple[ExtractedFragment, ...]:
    presentation = Presentation(BytesIO(content))
    fragments: list[ExtractedFragment] = []
    for number, slide in enumerate(presentation.slides, start=1):
        chunks: list[str] = []
        for shape in slide.shapes:
            chunks.extend(_shape_chunks(shape))
        fragments.append(ExtractedFragment(SLIDE, str(number), "\n".join(chunks)))
    return tuple(fragments)


def _iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _render_table(table: Table) -> str:
    lines: list[str] = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _render_rows(rows) -> str:
    lines: list[str] = []
    for row in rows:
        cells = [_cell_text(value) for value in row]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == value.microsecond == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decode_csv(content: bytes) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16")
    if content.startswith(b"\xef\xbb\xbf"):
        return content.decode("utf-8-sig")
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("cp1252")


# Semicolon before comma: Chilean spreadsheets exported as CSV use ";".
_DELIMITERS = (";", ",", "\t", "|")


def _detect_delimiter(text: str) -> str:
    lines = [line for line in text.splitlines() if line.strip()][:20]
    if not lines:
        return ","
    best: tuple[int, int, str] | None = None
    for delimiter in _DELIMITERS:
        counts = [line.count(delimiter) for line in lines]
        present = sum(count > 0 for count in counts)
        if present == 0:
            continue
        score = (present, sum(counts))
        if best is None or score > best[:2]:
            best = (*score, delimiter)
    return best[2] if best is not None else ","


def _shape_chunks(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        chunks: list[str] = []
        for child in shape.shapes:
            chunks.extend(_shape_chunks(child))
        return chunks
    if getattr(shape, "has_table", False) and shape.has_table:
        lines: list[str] = []
        for row in shape.table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
        return ["\n".join(lines)] if lines else []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        return [text] if text else []
    return []


_READERS[".pdf"] = extract_pdf
_READERS[".docx"] = extract_docx
_READERS[".xlsx"] = extract_xlsx
_READERS[".csv"] = extract_csv
_READERS[".pptx"] = extract_pptx
