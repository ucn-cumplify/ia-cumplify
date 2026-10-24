"""POST /api/v1/documents/extract returns the text of one knowledge-base document (step 2B.7)."""

import logging
import zipfile
from io import BytesIO

import pytest
from chat_http import HEADERS, SECRET, configure
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter
from test_evidence_text import _declare_uncompressed, _pdf, _save, _slides

from ia_cumplify.adapters.inbound.http.app import create_app
from ia_cumplify.adapters.inbound.http.routers import documents
from ia_cumplify.adapters.outbound.documents.text_extractor import extract_text
from ia_cumplify.application.use_cases import extract_document
from ia_cumplify.application.use_cases.extract_document import ExtractDocumentUseCase, utf16_length
from ia_cumplify.domain.document_extraction import (
    FRAGMENT_SEPARATOR,
    TEXT_TOO_LONG_DETAIL,
    UNREADABLE_FILE_DETAIL,
)
from ia_cumplify.domain.evidence_text import READ, ExtractedFile, ExtractedFragment

EXTRACT = "/api/v1/documents/extract"
KEYS = ["filename", "status", "detail", "text", "fragments"]


def test_a_pdf_comes_back_per_page_and_text_joins_the_pages(monkeypatch):
    response = _post(monkeypatch, "procedimiento.pdf", _pdf(["Entrega de EPP.", "", "Registro firmado."]))

    assert response.status_code == 200
    body = response.json()
    assert list(body) == KEYS
    assert body["filename"] == "procedimiento.pdf"
    assert body["status"] == "read"
    assert body["detail"] is None
    # The blank page is not a fragment; the others keep their page number.
    assert body["fragments"] == [
        {"locator": "page", "reference": "1", "text": "Entrega de EPP."},
        {"locator": "page", "reference": "3", "text": "Registro firmado."},
    ]
    assert body["text"] == "Entrega de EPP.\n\nRegistro firmado."
    _assert_text_is_the_joined_fragments(body)


def test_word_excel_csv_and_powerpoint_are_read_with_their_locators(monkeypatch):
    document = Document()
    document.add_paragraph("Uso obligatorio de casco.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Área"
    table.cell(0, 1).text = "Faena"
    workbook = Workbook()
    workbook.active.title = "Entregas"
    workbook.active.append(["Casco", 10])
    workbook.create_sheet("Bajas").append(["Guantes", 2])

    cases = {
        "manual.docx": (_save(document.save), [("body", "", "Uso obligatorio de casco.\nÁrea | Faena")]),
        "epp.xlsx": (_save(workbook.save), [("sheet", "Entregas", "Casco | 10"), ("sheet", "Bajas", "Guantes | 2")]),
        "epp.csv": ("Elemento;Cantidad\nCasco;10\n".encode(), [("body", "", "Elemento;Cantidad\nCasco;10")]),
        "charla.pptx": (_slides(2), [("slide", "1", "Diapositiva 1"), ("slide", "2", "Diapositiva 2")]),
    }
    for filename, (content, expected) in cases.items():
        body = _post(monkeypatch, filename, content).json()

        assert body["status"] == "read", filename
        assert [(item["locator"], item["reference"], item["text"]) for item in body["fragments"]] == expected
        _assert_text_is_the_joined_fragments(body)


def test_an_image_and_a_scanned_pdf_are_unreadable_and_carry_no_text(monkeypatch):
    for filename, content in (("foto.png", b"\x89PNG\r\n\x1a\n"), ("escaneo.pdf", _pdf(["", ""]))):
        body = _post(monkeypatch, filename, content).json()

        assert body == {"filename": filename, "status": "unreadable", "detail": None, "text": None, "fragments": []}


def test_an_unknown_extension_is_unsupported(monkeypatch):
    for filename in ("antiguo.doc", "notas.txt", "sin-extension"):
        body = _post(monkeypatch, filename, SECRET.encode()).json()

        assert body == {"filename": filename, "status": "unsupported", "detail": None, "text": None, "fragments": []}


def test_a_broken_an_encrypted_or_an_oversized_file_is_an_error_with_a_fixed_detail(monkeypatch):
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.encrypt("clave")
    encrypted = _save(writer.write)
    oversized = _declare_uncompressed(_save(Document().save), 65 * 1024 * 1024)
    cases = {
        "roto.pdf": b"%PDF-1.4 " + SECRET.encode(),
        "vacio.docx": b"",
        "cifrado.pdf": encrypted,
        "bomba.docx": oversized,
        "largo.pdf": _pdf([f"Pagina {number}" for number in range(101)]),
    }
    for filename, content in cases.items():
        body = _post(monkeypatch, filename, content).json()

        assert body == {
            "filename": filename,
            "status": "error",
            "detail": UNREADABLE_FILE_DETAIL,
            "text": None,
            "fragments": [],
        }, filename


def test_a_parser_message_that_quotes_the_file_does_not_leave_the_service(monkeypatch, caplog):
    # lxml names the undefined entity: the extractor's own detail would carry it.
    leaking = _docx_with_body(f"<w:p><w:r><w:t>Sueldo &{SECRET};</w:t></w:r></w:p>")
    assert SECRET in (extract_text(".docx", leaking).detail or "")
    caplog.set_level(logging.DEBUG)

    response = _post(monkeypatch, "planilla.docx", leaking)

    assert response.json()["status"] == "error"
    assert response.json()["detail"] == UNREADABLE_FILE_DETAIL
    assert SECRET not in response.text
    assert SECRET not in caplog.text


def test_text_past_the_cap_is_an_error_and_never_comes_back_cut(monkeypatch):
    rows = "".join(f"Fila {number};" + "x" * 50 + "\n" for number in range(5_000))

    body = _post(monkeypatch, "registro.csv", rows.encode()).json()

    assert body == {
        "filename": "registro.csv",
        "status": "error",
        "detail": TEXT_TOO_LONG_DETAIL,
        "text": None,
        "fragments": [],
    }


def test_the_cap_is_counted_in_utf16_units_as_the_backend_does(monkeypatch):
    # 110,000 characters outside the BMP: under the extractor's cap, 220,000 UTF-16 units. Short rows:
    # the csv module refuses a field past 131,072 characters.
    emoji = "\n".join(["\U0001f600" * 55] * 2_000)
    assert len(emoji) < 200_000 < utf16_length(emoji)

    body = _post(monkeypatch, "emoji.csv", emoji.encode()).json()

    assert body["status"] == "error"
    assert body["detail"] == TEXT_TOO_LONG_DETAIL
    assert body["text"] is None


def test_the_separators_count_toward_the_cap(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["a" * 10])
    workbook.create_sheet("Dos").append(["b" * 10])
    content = _save(workbook.save)

    monkeypatch.setattr(extract_document, "MAX_TEXT_UTF16_UNITS", 22)
    fits = ExtractDocumentUseCase().execute("base.xlsx", content)
    monkeypatch.setattr(extract_document, "MAX_TEXT_UTF16_UNITS", 21)
    passes = ExtractDocumentUseCase().execute("base.xlsx", content)

    assert fits.status == READ
    assert fits.text == "a" * 10 + FRAGMENT_SEPARATOR + "b" * 10
    assert passes.status == "error"
    assert passes.detail == TEXT_TOO_LONG_DETAIL
    assert passes.text is None
    assert passes.fragments == ()


def test_a_lone_surrogate_becomes_a_replacement_character(monkeypatch):
    def fake_extract_text(extension, content, limits):
        return ExtractedFile(READ, (ExtractedFragment("sheet", "Hoja\ud800", "Texto \udfff roto"),))

    monkeypatch.setattr(extract_document, "extract_text", fake_extract_text)

    body = _post(monkeypatch, "base.xlsx", b"x").json()

    assert body["fragments"] == [{"locator": "sheet", "reference": "Hoja\ufffd", "text": "Texto \ufffd roto"}]
    assert body["text"] == "Texto \ufffd roto"


def test_a_path_in_the_filename_is_reduced_to_its_last_segment(monkeypatch):
    body = _post(monkeypatch, r"C:\fakepath\acta.csv", b"Acta firmada").json()

    assert body["filename"] == "acta.csv"
    assert body["status"] == "read"


@pytest.mark.parametrize(
    "files",
    [
        [],
        [("file", ("a.csv", b"uno")), ("file", ("b.csv", b"dos"))],
        [("file", ("C:\\fakepath\\", b"uno"))],
        [("file", ("a" * 252 + ".csv", b"uno"))],
        [("other", ("a.csv", b"uno"))],
    ],
    ids=["no-file", "two-files", "no-name", "long-name", "other-part"],
)
def test_anything_but_one_named_file_is_422(monkeypatch, files):
    configure(monkeypatch)

    with TestClient(create_app()) as client:
        response = client.post(EXTRACT, headers=HEADERS, files=files or None, data={"note": SECRET})

    assert response.status_code == 422
    assert SECRET not in response.text


def test_a_file_past_the_size_cap_is_422_without_its_content(monkeypatch):
    monkeypatch.setattr(documents, "MAX_FILE_BYTES", 1024 * 1024)

    response = _post(monkeypatch, "grande.csv", SECRET.encode() * (1024 * 1024 // len(SECRET) + 1))

    assert response.status_code == 422
    assert response.json() == {"detail": "The file is larger than 1 MB."}
    assert SECRET not in response.text


def test_the_log_carries_neither_the_filename_nor_the_text(monkeypatch, caplog):
    caplog.set_level(logging.DEBUG)

    response = _post(monkeypatch, f"{SECRET}.csv", f"Texto {SECRET}".encode())

    assert SECRET in response.json()["text"]
    assert "Document extraction:" in caplog.text
    assert SECRET not in caplog.text


def test_the_key_is_required_and_no_openai_key_is(monkeypatch):
    configure(monkeypatch)
    files = {"file": ("acta.csv", b"Acta")}

    with TestClient(create_app()) as client:
        missing = client.post(EXTRACT, files=files)
        wrong = client.post(EXTRACT, headers={"X-API-Key": "otra"}, files=files)
        right = client.post(EXTRACT, headers=HEADERS, files=files)

    assert missing.status_code == 401
    assert wrong.status_code == 401
    # conftest leaves OPENAI_API_KEY unset: extraction never calls the model.
    assert right.status_code == 200


def test_without_a_service_key_the_endpoint_is_503(monkeypatch):
    with TestClient(create_app()) as client:
        response = client.post(EXTRACT, headers=HEADERS, files={"file": ("acta.csv", SECRET.encode())})

    assert response.status_code == 503
    assert SECRET not in response.text


def _post(monkeypatch: pytest.MonkeyPatch, filename: str, content: bytes):
    configure(monkeypatch)
    with TestClient(create_app()) as client:
        return client.post(EXTRACT, headers=HEADERS, files={"file": (filename, content)})


def _assert_text_is_the_joined_fragments(body: dict) -> None:
    assert body["text"] == FRAGMENT_SEPARATOR.join(item["text"] for item in body["fragments"])
    assert all(item["text"].strip() for item in body["fragments"])


def _docx_with_body(xml: str) -> bytes:
    source = zipfile.ZipFile(BytesIO(_save(Document().save)))
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                data = data.replace(b"<w:body>", b"<w:body>" + xml.encode())
            package.writestr(info, data)
    return buffer.getvalue()
