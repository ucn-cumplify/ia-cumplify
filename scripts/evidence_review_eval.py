"""Manual evaluation set of evidence-review-v3: EVR-012 to EVR-031, EVR-034 and EVR-035 of docs/EvidenceReview/test.csv.

It spends real tokens. It builds each case's files and posts them to an ia-cumplify that is already
running, and that service calls OpenAI. It is not part of pytest, which collects only tests/.

    uv run python scripts/evidence_review_eval.py [--base-url URL] [--case ID ...] [--out FILE]
                                                 [--label low|medium] [--timeout SECONDS]
    uv run python scripts/evidence_review_eval.py --check-set
    uv run python scripts/evidence_review_eval.py --compare LOW.jsonl MEDIUM.jsonl

SERVICE_API_KEY comes from the environment or --api-key, never from .env. The cases carry their own
text and files: this script does not read the database. Responses go to --out, outside the repository.

The effort is the service's EVIDENCE_REVIEW_REASONING_EFFORT (empty uses OPENAI_REASONING_EFFORT).
Run the set once with low and once with medium, passing --label so --compare can put the two files
side by side. The plan estimates about US$0.20 per full run.
"""

from __future__ import annotations

import argparse
import http.client
import json
import os
import re
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
DEFAULT_SET = HERE / "evidence_review_eval_set.json"
NATURES = ("explicit_obligations", "informative", "referral", "procedural")
STATUSES = ("accredited", "partial", "not_accredited", "not_evaluable")
FILE_STATUSES = ("read", "unreadable", "unsupported", "error")
VERSION_PREFIX = "evidence-review-v3@"
# US$ per million tokens, the rates of the sprint 3 plan.
INPUT_PRICE = 0.20
OUTPUT_PRICE = 1.20
# Phrases that turn "what is missing" into a recommended action.
ACTION = re.compile(
    r"se recomienda|recomendamos|hay que|conviene|plan de acci[oó]n|la empresa debe|tienen que",
    re.IGNORECASE,
)
# A heuristic for "the missing text lists this period as missing": the word comes after a comma, «y», «e», «ni», «a»
# or «hasta», maybe followed by «de», as in «septiembre ni octubre» or «de abril a octubre», and not before words
# that say it is still running. «al 20 de octubre», the reference date, is not a listing.
LISTED_BEFORE = re.compile(r"(?:,\s*|\b(?:y|e|ni|a|hasta)\s+)(?:de\s+)?$", re.IGNORECASE)
RUNNING_AFTER = re.compile(r"^\s*(?:de\s+\d{4}\s+)?(?:aún|todavía|no\b|está|sigue|en curso)", re.IGNORECASE)


class SetError(ValueError):
    """The evaluation set does not have the expected shape."""


def load_set(path: Path) -> list[dict]:
    cases = json.loads(path.read_text(encoding="utf-8"))["cases"]
    seen: set[str] = set()
    for case in cases:
        for name in ("id", "title", "reference_date", "article", "legal_body", "files", "expected"):
            if name not in case:
                raise SetError(f"case {case.get('id')}: missing {name}")
        if case["id"] in seen:
            raise SetError(f"case {case['id']}: repeated id")
        seen.add(case["id"])
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", case["reference_date"]):
            raise SetError(f"case {case['id']}: reference_date must be YYYY-MM-DD")
        expected = case["expected"]
        if expected.get("nature") not in NATURES:
            raise SetError(f"case {case['id']}: nature must be one of {', '.join(NATURES)}")
        if expected.get("verdict") not in STATUSES:
            raise SetError(f"case {case['id']}: verdict must be one of {', '.join(STATUSES)}")
        for key in ("all_status", "has_status"):
            if key in expected and expected[key] not in STATUSES:
                raise SetError(f"case {case['id']}: {key} must be one of {', '.join(STATUSES)}")
        if expected.get("requirements") not in (None, "empty"):
            raise SetError(f"case {case['id']}: requirements must be empty or omitted")
        for key in ("missing_names", "missing_not_listed"):
            terms = expected.get(key, [])
            if not isinstance(terms, list) or not all(isinstance(term, str) and term.strip() for term in terms):
                raise SetError(f"case {case['id']}: {key} must be a list of words")
    return cases


def review_body(case: dict) -> dict:
    body = {
        "reference_date": case["reference_date"],
        "article": case["article"],
        "legal_body": case["legal_body"],
        "obligation": case.get("obligation"),
        "files": [{"id": item["id"], "filename": item["filename"]} for item in case["files"]],
    }
    return body


def build_files(case: dict) -> list[tuple[str, bytes]]:
    return [(item["filename"], _file_bytes(item)) for item in case["files"]]


def _file_bytes(spec: dict) -> bytes:
    if "pdf" in spec:
        pages = list(spec["pdf"])
        if spec.get("pad_to_chars"):
            pages.append("x" * int(spec["pad_to_chars"]))
        return _pdf(pages)
    if "docx" in spec:
        return _docx(spec["docx"])
    if "xlsx" in spec:
        return _xlsx(spec["xlsx"]["sheets"])
    if "csv" in spec:
        return spec["csv"].encode("utf-8")
    if "pptx" in spec:
        return _pptx(spec["pptx"])
    if spec.get("png"):
        return b"\x89PNG\r\n\x1a\nnot-parsed"
    raise SetError(f"file {spec.get('filename')}: no content")


def _pdf(pages: list[str]) -> bytes:
    """A text PDF. Helvetica is WinAnsi, so the page text is Latin-1."""
    streams = []
    for text in pages:
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        streams.append(f"BT /F1 12 Tf 72 700 Td ({escaped}) Tj ET".encode("latin-1"))
    objects: dict[int, bytes] = {}
    kids = " ".join(f"{4 + index * 2} 0 R" for index in range(len(streams)))
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(streams)} >>".encode()
    objects[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    for index, stream in enumerate(streams):
        page_id = 4 + index * 2
        content_id = page_id + 1
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        ).encode()
        objects[content_id] = b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    chunks = [b"%PDF-1.4\n"]
    offsets = [0]
    for number in range(1, max(objects) + 1):
        offsets.append(sum(len(chunk) for chunk in chunks))
        chunks.append(f"{number} 0 obj\n".encode() + objects[number] + b"\nendobj\n")
    xref_at = sum(len(chunk) for chunk in chunks)
    size = max(objects) + 1
    xref = [f"xref\n0 {size}\n".encode(), b"0000000000 65535 f \n"]
    xref.extend(f"{offsets[number]:010d} 00000 n \n".encode() for number in range(1, size))
    trailer = f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return b"".join([*chunks, *xref, trailer])


def _docx(paragraphs: list[str]) -> bytes:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    return _save(document.save)


def _xlsx(sheets: list[dict]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    for index, sheet in enumerate(sheets):
        current = workbook.active if index == 0 else workbook.create_sheet()
        current.title = sheet["name"]
        for row in sheet["rows"]:
            current.append(row)
    return _save(workbook.save)


def _pptx(slides: list[str]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for text in slides:
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        box.text_frame.text = text
    return _save(presentation.save)


def _save(write) -> bytes:
    buffer = BytesIO()
    write(buffer)
    return buffer.getvalue()


def local_files(case: dict) -> list[str]:
    """What extract_text says about the built files, before any request is sent."""
    from ia_cumplify.adapters.outbound.documents.text_extractor import extract_text

    problems = []
    expected = {item["filename"]: item for item in case["expected"].get("files", [])}
    for filename, content in build_files(case):
        dot = filename.rfind(".")
        reading = extract_text(filename[dot:], content)
        wanted = expected.get(filename)
        if wanted is None:
            continue
        if reading.status != wanted["status"] or reading.truncated is not wanted["truncated"]:
            problems.append(
                f"{filename}: got {reading.status} truncated={reading.truncated}, "
                f"expected {wanted['status']} truncated={wanted['truncated']}"
            )
    return problems


def checks(case: dict, body: dict) -> list[str]:
    """Differences between the response and what the case expects. Empty means it passed."""
    expected = case["expected"]
    found: list[str] = []
    if body.get("nature") != expected["nature"]:
        found.append(f"nature {body.get('nature')!r}, expected {expected['nature']}")
    if body.get("verdict") != expected["verdict"]:
        found.append(f"verdict {body.get('verdict')!r}, expected {expected['verdict']}")
    version = body.get("review_version") or ""
    if not version.startswith(VERSION_PREFIX):
        found.append(f"review_version {version!r}")
    requirements = body.get("requirements") or []
    if expected.get("requirements") == "empty" and requirements:
        found.append(f"{len(requirements)} requirements, expected none")
    if "all_status" in expected and any(item.get("status") != expected["all_status"] for item in requirements):
        found.append(f"a requirement is not {expected['all_status']}")
    if "has_status" in expected and all(item.get("status") != expected["has_status"] for item in requirements):
        found.append(f"no requirement is {expected['has_status']}")
    if expected.get("missing") == "empty" and any((item.get("missing") or "").strip() for item in requirements):
        found.append("missing is not empty")
    for item in requirements:
        if ACTION.search(item.get("missing") or ""):
            found.append("missing recommends an action")
            break
    missing = " ".join(item.get("missing") or "" for item in requirements)
    for term in expected.get("missing_names", []):
        if not re.search(rf"\b{re.escape(term)}\b", missing, re.IGNORECASE):
            found.append(f"missing does not name {term}")
    for term in expected.get("missing_not_listed", []):
        if listed(term, missing):
            found.append(f"missing lists {term}")
    known = {item["id"] for item in case["files"]}
    for item in requirements:
        for citation in item.get("citations") or []:
            if citation.get("file_id") not in known or not str(citation.get("quote") or "").strip():
                found.append("a citation names an unknown file or a blank quote")
                break
    by_name = {item.get("filename"): item for item in body.get("files") or []}
    for wanted in expected.get("files", []):
        got = by_name.get(wanted["filename"])
        if got is None:
            found.append(f"response has no {wanted['filename']}")
            continue
        if got.get("status") != wanted["status"] or got.get("truncated") is not wanted["truncated"]:
            found.append(
                f"{wanted['filename']}: {got.get('status')} truncated={got.get('truncated')}, "
                f"expected {wanted['status']} truncated={wanted['truncated']}"
            )
    return found


def listed(term: str, text: str) -> bool:
    """Whether text lists term among what is missing, by LISTED_BEFORE and RUNNING_AFTER."""
    for match in re.finditer(rf"\b{re.escape(term)}\b", text, re.IGNORECASE):
        if LISTED_BEFORE.search(text[: match.start()]) and not RUNNING_AFTER.match(text[match.end() :]):
            return True
    return False


def encode_multipart(review: str, files: list[tuple[str, bytes]]) -> tuple[bytes, str]:
    boundary = "----evidence-review-eval"
    chunks: list[bytes] = []
    chunks.append(f"--{boundary}\r\n".encode())
    chunks.append(b'Content-Disposition: form-data; name="review"\r\n\r\n')
    chunks.append(review.encode("utf-8"))
    chunks.append(b"\r\n")
    for filename, content in files:
        chunks.append(f"--{boundary}\r\n".encode())
        header = f'Content-Disposition: form-data; name="files"; filename="{filename}"\r\n'
        chunks.append(header.encode())
        chunks.append(b"Content-Type: application/octet-stream\r\n\r\n")
        chunks.append(content)
        chunks.append(b"\r\n")
    chunks.append(f"--{boundary}--\r\n".encode())
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


@dataclass
class Result:
    case: dict
    label: str
    status: int = 0
    body: dict | None = None
    detail: str = ""
    checks: list[str] = field(default_factory=list)
    total_ms: float = 0.0

    @property
    def passed(self) -> bool:
        return self.status == 200 and not self.checks

    def cost(self, input_price: float, output_price: float) -> float:
        usage = (self.body or {}).get("usage") or {}
        return (usage.get("prompt_tokens", 0) * input_price + usage.get("completion_tokens", 0) * output_price) / 1_000_000


def post_case(base_url: str, api_key: str, case: dict, timeout: float) -> tuple[int, dict | str, float]:
    import time

    review = json.dumps(review_body(case), ensure_ascii=False)
    payload, content_type = encode_multipart(review, build_files(case))
    url = urlsplit(base_url)
    path = (url.path.rstrip("/") or "") + "/api/v1/evidence-reviews"
    connection_class = http.client.HTTPSConnection if url.scheme == "https" else http.client.HTTPConnection
    started = time.perf_counter()
    connection = connection_class(url.hostname, url.port or (443 if url.scheme == "https" else 80), timeout=timeout)
    try:
        connection.request(
            "POST",
            path,
            body=payload,
            headers={
                "X-API-Key": api_key,
                "Content-Type": content_type,
                "Content-Length": str(len(payload)),
            },
        )
        response = connection.getresponse()
        raw = response.read()
    finally:
        connection.close()
    elapsed = (time.perf_counter() - started) * 1000
    try:
        parsed = json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        parsed = raw.decode("utf-8", errors="replace")[:500]
    return response.status, parsed, elapsed


def write_out(path: Path, results: list[Result], input_price: float, output_price: float) -> None:
    with path.open("w", encoding="utf-8") as out:
        for result in results:
            out.write(
                json.dumps(
                    {
                        "id": result.case["id"],
                        "title": result.case["title"],
                        "label": result.label,
                        "status": result.status,
                        "passed": result.passed,
                        "checks": result.checks,
                        "detail": result.detail,
                        "response": result.body,
                        "total_ms": round(result.total_ms, 1),
                        "cost_usd": round(result.cost(input_price, output_price), 6) if result.body else None,
                    },
                    ensure_ascii=False,
                )
                + "\n"
            )


def summarize(results: list[Result], input_price: float, output_price: float) -> bool:
    passed = 0
    for result in results:
        mark = "ok" if result.passed else "FAIL"
        extra = "" if result.passed else ": " + "; ".join(result.checks or [result.detail or str(result.status)])
        print(f"{mark} {result.case['id']} {result.case['title']} ({result.total_ms:.0f} ms){extra}")
        passed += result.passed
    answered = [result for result in results if result.body]
    prompt = sum((result.body or {}).get("usage", {}).get("prompt_tokens", 0) for result in answered)
    completion = sum((result.body or {}).get("usage", {}).get("completion_tokens", 0) for result in answered)
    cost = sum(result.cost(input_price, output_price) for result in answered)
    print(
        f"{passed}/{len(results)} passed. Tokens: {prompt} in, {completion} out. "
        f"Estimated cost US${cost:.4f} ({len(answered)} calls, US${input_price:.2f}/US${output_price:.2f} per million)."
    )
    return passed == len(results)


def compare(low_path: Path, medium_path: Path) -> int:
    low = _by_id(low_path)
    medium = _by_id(medium_path)
    ids = list(dict.fromkeys([*low, *medium]))
    same = 0
    print(f"{'id':<10} {'low':<6} {'medium':<8} {'verdicts':<28} {'ms low':>8} {'ms med':>8} {'US$ low':>8} {'US$ med':>8}")
    for case_id in ids:
        left = low.get(case_id)
        right = medium.get(case_id)
        left_verdict = (left or {}).get("response", {}) and (left or {}).get("response", {}).get("verdict")
        right_verdict = (right or {}).get("response", {}) and (right or {}).get("response", {}).get("verdict")
        if left and right and left.get("passed") and right.get("passed") and left_verdict == right_verdict:
            same += 1
        print(
            f"{case_id:<10} {_mark(left):<6} {_mark(right):<8} "
            f"{(left_verdict or '-')} / {(right_verdict or '-'):<16} "
            f"{_ms(left):>8} {_ms(right):>8} {_usd(left):>8} {_usd(right):>8}"
        )
    print(f"{same}/{len(ids)} cases passed on both runs with the same verdict.")
    return 0 if same == len(ids) and low and medium else 1


def _by_id(path: Path) -> dict[str, dict]:
    rows = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            row = json.loads(line)
            rows[row["id"]] = row
    return rows


def _mark(row: dict | None) -> str:
    if row is None:
        return "-"
    return "ok" if row.get("passed") else "FAIL"


def _ms(row: dict | None) -> str:
    return "-" if not row else f"{row.get('total_ms', 0):.0f}"


def _usd(row: dict | None) -> str:
    if not row or row.get("cost_usd") is None:
        return "-"
    return f"{row['cost_usd']:.4f}"


def default_out(label: str) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return Path(tempfile.gettempdir()) / f"evidence-review-eval-{label}-{stamp}.jsonl"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--api-key", default=os.environ.get("SERVICE_API_KEY", ""), help="defecto: SERVICE_API_KEY")
    parser.add_argument("--set", type=Path, default=DEFAULT_SET, dest="case_set")
    parser.add_argument("--case", action="append", default=[], help="solo estos casos (se puede repetir)")
    parser.add_argument("--label", default="low", choices=("low", "medium"), help="con qué esfuerzo corre el servicio")
    parser.add_argument("--out", type=Path, default=None, help="respuestas completas; defecto: el directorio temporal")
    parser.add_argument("--timeout", type=float, default=120.0, help="segundos por pedido; más que los 80 s de ia")
    parser.add_argument("--check-set", action="store_true", help="arma los archivos y no llama al servicio")
    parser.add_argument("--compare", nargs=2, type=Path, metavar=("LOW", "MEDIUM"))
    parser.add_argument("--input-price", type=float, default=INPUT_PRICE)
    parser.add_argument("--output-price", type=float, default=OUTPUT_PRICE)
    args = parser.parse_args(argv)

    if args.compare:
        return compare(args.compare[0], args.compare[1])

    try:
        cases = load_set(args.case_set)
    except (OSError, json.JSONDecodeError, SetError) as exc:
        parser.error(str(exc))
    if args.case:
        unknown = set(args.case) - {case["id"] for case in cases}
        if unknown:
            parser.error(f"casos desconocidos: {', '.join(sorted(unknown))}")
        cases = [case for case in cases if case["id"] in args.case]

    broken = [(case["id"], local_files(case)) for case in cases]
    broken = [(case_id, problems) for case_id, problems in broken if problems]
    if broken:
        for case_id, problems in broken:
            print(f"FAIL {case_id}: {'; '.join(problems)}", file=sys.stderr)
        return 1
    if args.check_set:
        print(f"{len(cases)} cases build the expected files.")
        return 0

    if not args.api_key:
        parser.error("falta SERVICE_API_KEY (o --api-key)")
    out = (args.out or default_out(args.label)).resolve()
    if out == REPO or REPO in out.parents:
        parser.error("--out tiene que quedar fuera del repositorio")

    results: list[Result] = []
    for case in cases:
        result = Result(case=case, label=args.label)
        try:
            result.status, parsed, result.total_ms = post_case(args.base_url, args.api_key, case, args.timeout)
        except (TimeoutError, OSError) as exc:
            result.detail = type(exc).__name__
            results.append(result)
            continue
        if isinstance(parsed, dict):
            result.body = parsed
            if result.status == 200:
                result.checks = checks(case, parsed)
            else:
                result.detail = str(parsed.get("detail", ""))[:300]
        else:
            result.detail = str(parsed)
        results.append(result)

    write_out(out, results, args.input_price, args.output_price)
    print(f"Respuestas: {out}")
    return 0 if summarize(results, args.input_price, args.output_price) else 1


if __name__ == "__main__":
    raise SystemExit(main())
